"""把 Markdown 格式的小说大纲转换为 Word (.docx) 文件。"""
import re

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

FONT = "微软雅黑"


def _set_east_asia(rpr, name=FONT):
    """为 rPr 设置中文字体（w:eastAsia）。"""
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    rfonts.set(qn("w:eastAsia"), name)


def _style_run(run, name=FONT):
    """给一个 run 设置中文字体，保证中文正常显示。"""
    run.font.name = name
    _set_east_asia(run._element.get_or_add_rPr(), name)


def _add_inline_runs(paragraph, text):
    """解析 **加粗**、*斜体*、`代码` 等行内 Markdown，逐段写入。"""
    tokens = re.split(r"(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`)", text)
    for tok in tokens:
        if not tok:
            continue
        if tok.startswith("**") and tok.endswith("**") and len(tok) > 4:
            run = paragraph.add_run(tok[2:-2])
            run.bold = True
            _style_run(run)
        elif tok.startswith("*") and tok.endswith("*") and len(tok) > 2:
            run = paragraph.add_run(tok[1:-1])
            run.italic = True
            _style_run(run)
        elif tok.startswith("`") and tok.endswith("`") and len(tok) > 2:
            run = paragraph.add_run(tok[1:-1])
            run.font.name = "Consolas"
            _style_run(run)
        else:
            _style_run(paragraph.add_run(tok))


def markdown_to_docx(md_text, title="小说大纲"):
    """把 Markdown 文本转换成 Document 对象。"""
    doc = Document()

    # 默认字体（正文，含中文）
    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(11)
    _set_east_asia(normal.element.get_or_add_rPr(), FONT)

    # 文档标题
    title_p = doc.add_heading(title, level=0)
    for run in title_p.runs:
        _style_run(run)

    lines = (md_text or "").replace("\r\n", "\n").split("\n")
    in_code = False
    code_buf = []

    def flush_code():
        p = doc.add_paragraph()
        run = p.add_run("\n".join(code_buf))
        run.font.name = "Consolas"
        _style_run(run)
        del code_buf[:]

    for line in lines:
        stripped = line.strip()

        if stripped.startswith("```"):
            if in_code:
                flush_code()
                in_code = False
            else:
                in_code = True
            continue
        if in_code:
            code_buf.append(line)
            continue

        if not stripped:
            continue

        # 标题
        m = re.match(r"^(#{1,6})\s+(.*)$", stripped)
        if m:
            level = len(m.group(1))
            h = doc.add_heading(m.group(2), level=level)
            for run in h.runs:
                _style_run(run)
                run.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
            continue

        # 分隔线：跳过
        if re.match(r"^(-{3,}|\*{3,})$", stripped):
            continue

        # 无序列表
        if re.match(r"^[-*+]\s+", stripped):
            p = doc.add_paragraph(style="List Bullet")
            _add_inline_runs(p, re.sub(r"^[-*+]\s+", "", stripped))
            continue

        # 有序列表
        if re.match(r"^\d+[.)]\s+", stripped):
            p = doc.add_paragraph(style="List Number")
            _add_inline_runs(p, re.sub(r"^\d+[.)]\s+", "", stripped))
            continue

        # 引用
        if re.match(r"^>\s?", stripped):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.3)
            _add_inline_runs(p, re.sub(r"^>\s?", "", stripped))
            continue

        # 普通段落
        _add_inline_runs(doc.add_paragraph(), stripped)

    if code_buf:
        flush_code()

    return doc
